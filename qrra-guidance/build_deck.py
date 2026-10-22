import sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

SRC, OUT = sys.argv[1], sys.argv[2]

TEAL = RGBColor(0x00, 0x30, 0x3C)
NAVY = RGBColor(0x09, 0x19, 0x20)
LIME = RGBColor(0xE7, 0xFF, 0x00)
PALE = RGBColor(0xEC, 0xFF, 0xC0)
GREEN = RGBColor(0xBF, 0xD8, 0x33)
STONE = RGBColor(0xE2, 0xDE, 0xD9)
MIST = RGBColor(0xF2, 0xF4, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x4A, 0x5A, 0x60)
FONT = "Arial"
HEAD = "Arial Black"

prs = Presentation(SRC)
L = {l.name: l for l in prs.slide_layouts}

# drop the template's sample slides
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)


def ph(slide, idx):
    for s in slide.placeholders:
        if s.placeholder_format.idx == idx:
            return s
    raise KeyError(idx)


def drop_ph(slide, idx):
    s = ph(slide, idx)
    s._element.getparent().remove(s._element)


def set_text(shape, text, size=None, bold=None, color=None):
    tf = shape.text_frame
    lines = text if isinstance(text, list) else [text]
    tf.text = ""
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        if size: r.font.size = Pt(size)
        if bold is not None: r.font.bold = bold
        if color is not None: r.font.color.rgb = color


def runs_para(p, parts, size, color, font=FONT):
    for part in parts:
        if isinstance(part, str):
            part = (part, {})
        t, o = part
        r = p.add_run()
        r.text = t
        r.font.size = Pt(o.get("size", size))
        r.font.name = o.get("font", font)
        r.font.bold = o.get("bold", False)
        r.font.italic = o.get("italic", False)
        r.font.color.rgb = o.get("color", color)


def text(slide, x, y, w, h, paras, size=14, color=TEAL, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, bullets=False, space=6, margin=0, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name: tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    fill_tf(tf, paras, size, color, bold, align, font, bullets, space)
    return tb


def fill_tf(tf, paras, size, color, bold=False, align=PP_ALIGN.LEFT, font=FONT, bullets=False, space=6):
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        if isinstance(para, str):
            para = [(para, {"bold": bold})]
        runs_para(p, para, size, color, font)
        if bullets:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.22))))
            pPr.set("indent", str(-int(Inches(0.22))))
            bf = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
            bc = pPr.makeelement(qn("a:buChar"), {"char": "•"})
            pPr.append(bf); pPr.append(bc)


def box(slide, x, y, w, h, fill, paras=None, size=14, color=TEAL, bold=False, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, shape=MSO_SHAPE.ROUNDED_RECTANGLE, margin=0.15, line=None,
        bullets=False, space=4, font=FONT, name=None, radius=0.06):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name: s.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(1.25)
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    if paras is not None:
        fill_tf(tf, paras, size, color, bold, align, font, bullets, space)
    return s


def badge(slide, x, y, d, label, fill=LIME, color=TEAL, size=16):
    return box(slide, x, y, d, d, fill, [label], size=size, color=color, bold=True, align=PP_ALIGN.CENTER,
               anchor=MSO_ANCHOR.MIDDLE, shape=MSO_SHAPE.OVAL, margin=0, font=HEAD)


def table(slide, x, y, w, col_w, rows, size=12, head_size=12, row_h=0.36, head_fill=TEAL, band=True,
          first_bold=True):
    nrows, ncols = len(rows), len(rows[0])
    gs = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
    tbl = gs.table
    tblPr = tbl._tbl.tblPr
    # strip the default table style banding so our fills rule
    tblPr.set("bandRow", "0"); tblPr.set("firstRow", "0")
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for r in range(nrows):
        tbl.rows[r].height = Inches(row_h)
        for c in range(ncols):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            val = rows[r][c]
            tf = cell.text_frame
            tf.word_wrap = True
            if r == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = head_fill
                fill_tf(tf, val, head_size, WHITE, True, space=0)
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = MIST if (band and r % 2 == 0) else WHITE
                fill_tf(tf, val, size, TEAL, first_bold and c == 0, space=0)
    return gs


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def content(layout, title):
    s = prs.slides.add_slide(L[layout])
    set_text(ph(s, 13), title)
    drop_ph(s, 14)
    return s


def section(num, title, sub):
    s = prs.slides.add_slide(L["Section Header"])
    set_text(s.shapes.title, title)
    set_text(ph(s, 1), sub)
    text(s, 0.48, 1.45, 3, 0.6, [[(num, {"font": HEAD, "color": LIME})]], size=28)
    return s


# ---------------------------------------------------------------- 1 Title
s = prs.slides.add_slide(L["Title Slide"])
set_text(s.shapes.title, "Conducting a QRRA")
set_text(ph(s, 1), ["Guidance for the Quality Department on facilitating Quality Requirements Risk Assessments",
                    "SLP 2.3.2.7_02  |  SLF 2.3.2.7_02"])
sub = ph(s, 1).text_frame.paragraphs
for r in sub[1].runs:
    r.font.size = Pt(20); r.font.color.rgb = LIME
notes(s, "Welcome. This session is for colleagues across the Quality Department. It covers how a Quality "
         "Requirements Risk Assessment (QRRA) is conducted under SLP 2.3.2.7_02, and, most importantly, the role "
         "we play in it. The key message: the QRRA is led and initiated by the Demander - the person accountable "
         "for specifying and buying the item. Quality facilitates and acts as the intelligent client.")

# ---------------------------------------------------------------- 2 Purpose
s = content("Comparison", "Why we are here")
objs = [("1", "Ownership", "Understand who leads the QRRA and where it fits in specification and procurement"),
        ("2", "Facilitation", "Know how to support a QRRA session as the intelligent client, not the author"),
        ("3", "Consistency", "Apply SLF 2.3.2.7_02 the same way every time: risk, Quality Grade, requirements, sign-off")]
y = 1.3
for n, h, b in objs:
    badge(s, 0.48, y + 0.12, 0.7, n)
    text(s, 1.4, y, 6.0, 1.4, [[(h, {"bold": True, "size": 20})], b], size=16, space=4)
    y += 1.75
box(s, 7.9, 1.3, 4.95, 5.1, TEAL, [
    [("What has changed", {"bold": True, "size": 20, "color": LIME})],
    [("SLP 1.02.45", {"bold": True}), " is now ", ("SLP 2.3.2.7_02", {"bold": True, "color": LIME}),
     " (Issue 1, effective 04/2026)"],
    [("SLF 1.02.45.01", {"bold": True}), " is now ", ("SLF 2.3.2.7_02", {"bold": True, "color": LIME}),
     " - the QRRA form"],
    [("Content has transferred across; the process is unchanged in intent.", {})],
    [("Watch out: ", {"bold": True, "color": LIME}),
     "SLP 4.06.08 (Issue 14) and older specifications still cite SLP 1.02.45. Read these as SLP 2.3.2.7_02."],
], size=16, color=WHITE, margin=0.35, space=12, radius=0.04)
notes(s, "Three objectives for today. Also flag the document change: SLP 1.02.45 has been superseded by SLP "
         "2.3.2.7_02 and the form SLF 1.02.45.01 by SLF 2.3.2.7_02. The amendment record confirms content was "
         "transferred from SLP 1.02.45. Some documents, including SLP 4.06.08 Issue 14, still reference the old "
         "number - when you see it, point people to SLP 2.3.2.7_02 and the current form.")

# ---------------------------------------------------------------- 3 Section
s = section("01", "Who owns the QRRA?", "The Demander leads. Quality facilitates.")
notes(s, "First, ownership. This is the point most often misunderstood.")

# ---------------------------------------------------------------- 4 Roles
s = content("1_Comparison", "Roles in a QRRA")
cards = [
    ("LEADS", "Demander", LIME, TEAL,
     "Accountable for specification and procurement, e.g. Project Manager, Design Manager, PEM, Responsible "
     "Engineer or Demand Specifier. Initiates the QRRA, owns the risk assessment and records it on SLF 2.3.2.7_02."),
    ("NEEDS", "Requestor", STONE, TEAL,
     "The person wanting the goods or services. Checks for a material master, specification and quality "
     "requirements, then launches procurement once they are in place."),
    ("FACILITATES", "Quality (QM / QE)", TEAL, WHITE,
     "The intelligent client. Advises on Quality Grade, risk and documentation; challenges and endorses. "
     "Accountable for applying oversight post award (step 1.6)."),
    ("CONTRIBUTE", "Key stakeholders", STONE, TEAL,
     "ICG, Construction, Commissioning Start-up / Operations, Engineering & Maintenance, SMEs and, where "
     "appropriate, the supply chain."),
]
x = 0.48
for tag, name, fill, col, body in cards:
    box(s, x, 1.25, 2.9, 4.0, fill, [
        [(tag, {"bold": True, "size": 12, "color": col})],
        [(name, {"bold": True, "size": 20, "color": col})],
        [(body, {"color": col})]], size=14, margin=0.2, space=8)
    x += 3.09
box(s, 0.48, 5.5, 12.37, 1.0, MIST, [
    [("“It is essential that the Demander work collaboratively with key stakeholders”",
      {"italic": True, "size": 18}),
     ("   SLP 2.3.2.7_02, Purpose and Scope", {"size": 12, "color": GREY})]],
    anchor=MSO_ANCHOR.MIDDLE, margin=0.3)
notes(s, "SLP 2.3.2.7_02 defines the Demander as the accountable person responsible for product specification "
         "and procurement. SLP 4.06.08 says the same of the Demand Specifier: they ensure the risk assessment to "
         "assign the Quality Grade has been correctly carried out. So the QRRA belongs to them. The SLP also says "
         "the Demander must ensure the required oversight and assurance activities are captured on the QRRA. "
         "Quality is one of the key stakeholders, with a specific role: advise, challenge, endorse, and after "
         "contract award the QE is accountable for making sure the oversight is applied.")

# ---------------------------------------------------------------- 5 Intelligent client
s = prs.slides.add_slide(L["3_Comparison"])
set_text(ph(s, 13), "Quality as the intelligent client")
set_text(ph(s, 14), "Facilitate, challenge, endorse")
set_text(ph(s, 15), ["We bring the questions, the standards and the experience.",
                     "The Demander brings the requirement and owns the answer."])
drop_ph(s, 16)
do = ["Prompt with the Appendix 1 questions",
      "Challenge component granularity and consequence ratings",
      "Ask for a clear reason for each activity and LTR",
      "Advise on Quality Grade, CQRs and ICG involvement",
      "Share supplier, CFSI and lessons-learned insight",
      "Endorse the agreed QRRA as part of the review team"]
dont = ["Write the QRRA on the Demander’s behalf",
        "Set the Quality Grade in isolation",
        "Apply one grade generically to a whole system",
        "Accept “same as last time” without review",
        "Add blanket inspection or LTRs with no rationale",
        "Take over the procurement decision"]
for x, head, fill, items in ((4.6, "We do", PALE, do), (8.85, "We don’t", MIST, dont)):
    box(s, x, 0.55, 4.0, 5.9, fill, None)
    text(s, x + 0.3, 0.85, 3.4, 0.5, [[(head, {"bold": True, "size": 22})]], size=22)
    text(s, x + 0.3, 1.6, 3.4, 4.6, items, size=15, space=12, bullets=True)
notes(s, "This is what 'intelligent client' means in practice. We know enough about the requirement, the "
         "standards and the supply chain to ask the right questions and spot gaps, but we do not do the work for "
         "the Demander. If Quality writes the QRRA, ownership moves to the wrong place, the Demander's knowledge of "
         "how the item is used is lost, and the assessment becomes generic. The SLP is explicit that there should "
         "be a clear reason for each activity and that blanket gathering of records should be avoided.")

# ---------------------------------------------------------------- 6 Section
s = section("02", "Where the QRRA fits", "Triggers, process steps and the link to procurement")
notes(s, "Next, where the QRRA sits in the wider process.")

# ---------------------------------------------------------------- 7 Process
s = content("Comparison", "SLP 2.3.2.7_02 at a glance")
steps = [("1.1", "Understand what is needed and how it is used", "Demander"),
         ("1.2", "Develop the specification", "Demander"),
         ("1.3", "Identify and endorse quality requirements (QRRA)", "Demander"),
         ("1.4", "Create a material master (if new)", "Requestor"),
         ("1.5", "Launch procurement", "Requestor"),
         ("1.6", "Apply quality oversight and assurance", "Quality"),
         ("1.7", "Feed back on quality assurance and oversight", "Quality")]
w, gap, x0 = 1.65, 0.137, 0.48
for i, (n, t, who) in enumerate(steps):
    x = x0 + i * (w + gap)
    hl = n == "1.3"
    box(s, x, 1.3, w, 2.55, LIME if hl else MIST, [[(n, {"bold": True, "size": 20, "font": HEAD})], t],
        size=14, bold=hl, margin=0.14, space=6)
    box(s, x, 3.98, w, 0.45, TEAL if who != "Requestor" else STONE, [who], size=12, bold=True,
        color=WHITE if who != "Requestor" else TEAL, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.04)
text(s, 0.48, 4.7, 4, 0.4, [[("Where Quality comes in", {"bold": True})]], size=18)
touch = [("Steps 1.1 to 1.3", "Key stakeholder: advise and challenge as requirements are defined"),
         ("Step 1.3", "Integrated review team: endorse the QRRA and Quality Grade (and approve any QG 02P)"),
         ("Steps 1.6 to 1.7", "Accountable: oversight applied, CFSI and off-condition feedback")]
for i, (h, b) in enumerate(touch):
    box(s, 0.48 + i * 4.18, 5.2, 4.0, 1.25, PALE, [[(h, {"bold": True})], b], size=14, margin=0.18, space=4)
notes(s, "The flowchart on page 1 of SLP 2.3.2.7_02 has three swim lanes: Demander, Quality Manager (or nominated "
         "QE) and Requestor. The QRRA is produced at step 1.3, in the Demander's lane. Quality's own lane only "
         "starts at step 1.6 once procurement is launched. Before that we contribute as a key stakeholder and, at "
         "1.3, as part of the integrated review team that endorses the final Quality Grade. Note the loop back to "
         "1.3 if a CFSI issue is identified at step 1.6.")

# ---------------------------------------------------------------- 8 Triggers
s = content("1_Comparison", "When is a QRRA needed?")
cols = [("New plant and equipment",
         ["Triggered once safety and operational performance requirements are identified in the Product "
          "Verification Plan (SEMS 2.3.2.7)",
          "Verification activities for the supplier are taken from the PVP and set out as quality requirements"]),
        ("Existing assets, goods and services",
         ["Triggered by replacement: obsolete or faulty equipment, or any service requested from a supplier",
          "Evidence level comes from regulatory needs and Maintenance Schedules (SEMS 2.6.5.1)"]),
        ("Applies to everything",
         ["COTS, MOTS and bespoke items all complete step 1.3",
          "Previous supplier performance is no guarantee of future performance",
          "Review existing specifications and material masters: nothing should be assumed unchanged"])]
for i, (h, items) in enumerate(cols):
    box(s, 0.48 + i * 4.18, 1.3, 4.0, 3.55, MIST, [[(h, {"bold": True, "size": 18})]] + items,
        size=14, margin=0.25, space=8)
box(s, 0.48, 5.1, 12.37, 1.35, TEAL, [
    [("Exemptions and simplifications", {"bold": True, "color": LIME, "size": 16})],
    "QG 04 items bought from the electronic catalogue are exempt. Where only minimal oversight is needed (e.g. a "
    "calibration certificate) it can be recorded on the material master or purchase requisition."],
    size=14, color=WHITE, margin=0.25, space=4, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Triggers are in section 1.1 and 1.3 of the SLP. Stress that all items go through step 1.3, including "
         "COTS. Even repeat buys should be re-checked: SLP 2.3.2.7_02 says previous supplier performance is no "
         "guarantee and that where the Demander queries existing requirements they should contact Quality. The "
         "only exemption is QG 04 e-catalogue purchases. Technology-specific QRRA recommendations may also exist "
         "as addenda to Engineering Standards.")

# ---------------------------------------------------------------- 9 Procurement link
s = content("Comparison", "How the QRRA links to procurement (SLP 4.06.08)")
flow = [("1.1", "Make vs Buy analysis", "SLP 4.06.24"),
        ("1.2", "Specify the requirement and agree Quality Grade", "SLP 1.02.44 and SLP 2.3.2.7_02"),
        ("1.3", "Review requirements to select route", "Frameworks, e-catalogue"),
        ("1.4", "Raise Purchase Requisition", "Needs specification and Quality Grade"),
        ("1.8", "Procurement operations", "Supply Chain Directorate")]
w, gap = 2.27, 0.255
for i, (n, t, ref) in enumerate(flow):
    x = 0.48 + i * (w + gap)
    hl = n == "1.2"
    box(s, x, 1.3, w, 2.2, LIME if hl else MIST, [[(n, {"bold": True, "size": 18, "font": HEAD})],
                                                  [(t, {"bold": True})], [(ref, {"size": 12, "color": GREY})]],
        size=14, margin=0.16, space=6)
    if i < len(flow) - 1:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + w + 0.03), Inches(2.25), Inches(0.2), Inches(0.3))
        a.fill.solid(); a.fill.fore_color.rgb = TEAL; a.line.fill.background()
rules = [("Trained Demand Specifiers only", "The SAP Demander will not process a request from someone not on the "
          "trained Demand Specifier list"),
         ("No Purchase Cards for QG 01 to 03", "Purchase Cards must not be used for Quality Grade 01, 02 or 03 "
          "supplies or services"),
         ("QG 01 and 02 go to ICG", "The Supply Chain representative emails the Purchase Order and documents to ICG"),
         ("Check the QASL first", "No supplier without a grading on the Quality Assessed Supplier List (SLP 4.06.04)")]
for i, (h, b) in enumerate(rules):
    x = 0.48 + (i % 2) * 6.28
    y = 3.85 + (i // 2) * 1.35
    box(s, x, y, 6.09, 1.2, PALE, [[(h, {"bold": True})], b], size=14, margin=0.2, space=3, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "SLP 4.06.08 step 1.2 requires the Demand Specifier to specify the requirement and agree the Quality "
         "Grade, considering risk and mitigation in line with SLP 1.02.45 - now SLP 2.3.2.7_02. The Demand "
         "Specifier must also make sure the appropriate risk assessment is conducted to define the correct Quality "
         "Grade. The Quality Grade is a required input to the Purchase Requisition at step 1.4. Under SLP "
         "2.3.2.7_02 step 1.5, procurement can only commence once the specification, QRRA and lifetime records "
         "requirements are defined.")

# ---------------------------------------------------------------- 10 Section
s = section("03", "Running a QRRA", "Before, during and after the session")
notes(s, "Now the practical part: how a QRRA session runs.")

# ---------------------------------------------------------------- 11 Before
s = content("1_Comparison", "Before the session")
for x, head, tag, fill, col, items in (
        (0.48, "Demander prepares", "LEADS", LIME, TEAL,
         ["Defines the scope and how it will be broken down",
          "Brings the specification (SEMS 2.3.2.6 / SLP 1.02.44) or COTS datasheet with OEM part numbers",
          "Brings the PVP (SEMS 2.3.2.7) or Maintenance Schedule (SEMS 2.6.5.1)",
          "Gathers existing material masters and previous QRRAs",
          "Invites key stakeholders and books the session"]),
        (6.76, "Quality gets ready", "FACILITATES", TEAL, WHITE,
         ["Reviews the pack in advance and notes questions",
          "Selects relevant Appendix 1 prompts",
          "Checks supplier history on the QASL and any CFSI or lessons learned",
          "Identifies relevant standards and regulations (e.g. PUWER, LOLER, UKCA, BS EN 1090, BS EN 61508)",
          "Confirms ICG availability where QG 01 or 02 is likely"])):
    box(s, x, 1.3, 6.09, 5.15, fill, [[(tag, {"bold": True, "size": 12, "color": col})],
                                      [(head, {"bold": True, "size": 22, "color": col})]] + items,
        size=17, color=col, margin=0.35, space=14)
notes(s, "The Demander arranges and prepares the session. If a Demander asks Quality to 'do the QRRA', redirect "
         "them: offer to facilitate, and help them prepare. Our preparation is about asking good questions: "
         "Appendix 1 of SLP 2.3.2.7_02 is our main tool. Regulated items (PUWER, LOLER, UKCA) and items built to "
         "standards with defined oversight (BS EN 1090 execution class, BS EN 61508 SIL) must be included on the "
         "QRRA.")

# ---------------------------------------------------------------- 12 During
s = content("Comparison", "During the session: six steps through SLF 2.3.2.7_02")
six = [("Break down the scope", "Choose component granularity. Split Procurement / Fabrication from Site "
        "Installation. Include lifetime spares."),
       ("Assess consequences", "Rate each item against the seven risk columns on the form."),
       ("Agree the Quality Grade", "QG 01, 02, 02P, 03 or 04, driven by the consequences and the Appendix 2 "
        "‘other factors’."),
       ("Select requirements", "For each requirement, record who verifies it: SL, SL/MC, MC, Other or N/A."),
       ("Justify", "Record the reason for each activity and the factors considered in Supporting Comments."),
       ("Confirm", "Agree with all stakeholders and complete the ‘Confirmed by’ section.")]
for i, (h, b) in enumerate(six):
    x = 0.48 + (i % 3) * 4.18
    y = 1.3 + (i // 3) * 2.3
    box(s, x, y, 4.0, 2.1, MIST, None)
    badge(s, x + 0.22, y + 0.25, 0.6, str(i + 1))
    text(s, x + 1.0, y + 0.22, 2.8, 1.8, [[(h, {"bold": True, "size": 17})], b], size=14, space=4)
box(s, 0.48, 5.95, 12.37, 0.5, TEAL, [[("The Demander leads every step. ", {"bold": True, "color": LIME}),
                                       ("Quality prompts, challenges and checks the reasoning is recorded.",
                                        {"color": WHITE})]], size=14, anchor=MSO_ANCHOR.MIDDLE, margin=0.2)
notes(s, "These six steps follow the structure of the form. Step 1: the SLP says oversight should relate to "
         "specific equipment or components, not be applied generically to a system, and the Demander, Designer or "
         "Requestor decides the granularity with the stakeholders. Steps 2 to 4 are covered on the next three "
         "slides. Step 5: the SLP requires a clear reason for each activity with the applicable consequences "
         "recorded. Step 6: sign-off.")

# ---------------------------------------------------------------- 13 Risk columns
s = content("1_Comparison", "Step 2: rating the risk columns")
rows = [["Risk column", "Rating", "What to consider"],
        ["Nuclear / Radiological Safety", "SFC 1, SFC 2, SFC 3, SF, SM, SRE, NP, EE",
         "Categorisation per SLSP 1.02.14.02; exposure or loss of containment"],
        ["Conventional Safety", "VH, H, M, L, N/A", "COSHH, DSEAR, machinery, COMAH, LOLER, PUWER, pressure systems"],
        ["Environmental", "E1, E2, E3", "SLSP 1.02.14.02 s1.2, potential impact under fault conditions"],
        ["Operational (TRAM)", "VH, H, M, L, N/A", "Throughput, reliability, availability, maintainability"],
        ["Security", "VH, H, M, L, N/A", "Consequences of physical or cyber security failure"],
        ["Grey Market", "Is the OEM known?", "CFSI exposure (SEMS 3.3.5.3, IAEA NP-T-3.26)"],
        ["Business", "VH, H, M, L, N/A", "Cost, programme, commercial and reputational impact"]]
table(s, 0.48, 1.3, 12.37, [3.3, 3.6, 5.47], rows, size=14, head_size=14, row_h=0.6)
text(s, 0.48, 6.15, 12.37, 0.35, [[("Use Appendix 3 of SLP 2.3.2.7_02 for risk scoring criteria and consequences "
                                    "(Very High to Low).", {"italic": True, "color": GREY})]], size=12)
notes(s, "Ratings and guidance are from the Guidance tab of SLF 2.3.2.7_02. The Demander should rate these, using "
         "the Responsible Engineer and SMEs for nuclear safety categorisation. Quality's job is to test the "
         "ratings: what happens if this component fails at point of use? Is it accessible? Can it be tested after "
         "installation? Appendix 3 gives criteria and consequences for each risk level.")

# ---------------------------------------------------------------- 14 Quality Grade
s = content("Comparison", "Step 3: agreeing the Quality Grade")
rows = [["Grade", "Risk", "Risk mitigation", "Quality Plan"],
        ["QG 01", "Very High", "Inspection at supplier’s works by ICG", "Required, approved before manufacture"],
        ["QG 02", "High", "Inspection at supplier’s works by ICG", "Required, approved before manufacture"],
        ["QG 02P", "High (proprietary)", "Manufacturer inspection; final inspection at Lillyhall Stores",
         "Not required; justification on QRRA"],
        ["QG 03", "Medium", "Receipt inspection by ICG", "By exception"],
        ["QG 04", "Low / none", "Transit damage and order completion check only", "Not required"]]
table(s, 0.48, 1.3, 8.0, [1.05, 1.6, 3.15, 2.2], rows, size=13, head_size=13, row_h=0.78)
callouts = [("Minimum QG 03", "Any item needing design, development or Sellafield Ltd inspection"),
            ("QG 02P needs approval", "Justification and risk assessment recorded on the QRRA, approved by Quality "
             "and the integrated review team"),
            ("No agreement?", "The Demand Specifier seeks advice from Quality and / or the relevant SME")]
for i, (h, b) in enumerate(callouts):
    box(s, 8.75, 1.3 + i * 1.62, 4.1, 1.47, PALE if i < 2 else TEAL,
        [[(h, {"bold": True, "color": TEAL if i < 2 else LIME})], [(b, {"color": TEAL if i < 2 else WHITE})]],
        size=14, margin=0.2, space=3, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "From SLP 2.3.2.7_02 section 1.3 and Appendix 2. The grade sets which Contract Quality Requirements "
         "apply. Consider raising the grade if the item is complex, unique or novel, has a poor quality history, "
         "is inaccessible during operation, is hard to test after installation, or the potential economic loss is "
         "much larger than usual for that grade. The final grade is endorsed by the integrated review team, "
         "including the Requestor. Mandatory CQR deviations go through SEMS 3.3.6.2 to Supply Chain Quality.")

# ---------------------------------------------------------------- 15 Requirements
s = content("1_Comparison", "Step 4: selecting requirements and who verifies")
groups = [("Before manufacture", ["UKCA marking", "Start-up meeting", "Verification of contractor design",
                                  "Approval of Quality Plan / ITP"]),
          ("People and procedures", ["Nuclear Safety Tool Box", "Competent person qualifications",
                                     "Weld, paint, NDT or test procedures"]),
          ("During manufacture", ["Material control", "Interim inspection", "Specialist testing (PT, LT, CT, LK, CFSI)",
                                  "Acceptance test (FAT, SAT, CAT)"]),
          ("Release and records", ["Documentation (CofC, DofC, O&M, CC)", "Final inspection", "LTR review",
                                   "Release certificate (5059 / 5199)", "Packing and preservation"])]
for i, (h, items) in enumerate(groups):
    box(s, 0.48 + i * 3.13, 1.3, 2.95, 3.6, MIST, [[(h, {"bold": True, "size": 17})]] + items,
        size=14, margin=0.22, space=7)
codes = [("Yes (SL)", "Sellafield Ltd"), ("Yes (SL/MC)", "SL and Main Contractor"), ("Yes (MC)", "Main Contractor"),
         ("Yes (Other)", "As specified"), ("N/A", "Not applicable")]
text(s, 0.48, 5.1, 6, 0.35, [[("Verification codes: state who verifies each requirement", {"bold": True})]], size=14)
for i, (c, d) in enumerate(codes):
    box(s, 0.48 + i * 2.5, 5.55, 2.37, 0.9, TEAL, [[(c, {"bold": True, "color": LIME})], [(d, {"color": WHITE})]],
        size=13, margin=0.12, space=2, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "These are the requirement columns on SLF 2.3.2.7_02. The QRRA must identify which function delivers each "
         "element of assurance and oversight, e.g. Quality, ICG, Responsible Engineers, Construction or "
         "Commissioning. Extra requirements can be added as columns or in Supporting Comments. For large volumes, "
         "oversight can apply to first-off production and then revert to goods received inspection, defined in the "
         "ITP or Quality Plan. LTRs beyond those normally supplied with COTS go in the LTR schedule (SEMS "
         "1.2.10.1).")

# ---------------------------------------------------------------- 16 Prompts
s = prs.slides.add_slide(L["3_Comparison"])
set_text(ph(s, 13), "Questions to ask")
set_text(ph(s, 14), "From Appendix 1")
set_text(ph(s, 15), ["Use open questions to test the Demander’s thinking.",
                     "Let them answer and record it; don’t answer for them."])
drop_ph(s, 16)
qs = ["Is it clear what oversight and LTRs the COTS product already includes?",
      "Is it clear how these activities support the safety and security case?",
      "Is the item complex, unique or novel?",
      "Has the product or supplier had a poor-quality history?",
      "Is the item accessible during normal operations?",
      "Can functionality be proven after installation?",
      "Is material traceability needed to achieve the performance condition?",
      "Is in-process observation needed that final inspection can’t replace?",
      "What witnessing of functional testing is needed for confidence?",
      "Are the extra costs over the COTS standard understood and justified?"]
for c in range(2):
    box(s, 4.5 + c * 4.43, 0.55, 4.25, 5.9, MIST if c == 0 else PALE, qs[c * 5:(c + 1) * 5], size=15,
        margin=0.3, space=14, bullets=True)
notes(s, "A selection of the Appendix 1 questions. Appendix 1 has general factors, factors from IAEA GS-G-3.5, "
         "specific factors to be read with SLP 1.10.315, and CFSI / grey market factors. Bring a copy to every "
         "session. These questions keep the Demander in the lead: they own the answers, we make sure the "
         "questions have been asked.")

# ---------------------------------------------------------------- 17 CFSI
s = content("Comparison", "CFSI and grey market: what to probe")
flags = ["Price lower than expected",
         "Obsolete or hard-to-source items with short lead times",
         "Chain of custody: is the item traceable to the OEM via authorised stockists?",
         "Supplier is not the OEM, or is an independent distributor of unknown reputation"]
box(s, 0.48, 1.3, 6.09, 5.15, MIST, [[("Warning signs", {"bold": True, "size": 22})]] + flags,
    size=17, margin=0.35, space=16)
acts = [[("Record extra checks", {"bold": True, "color": LIME}),
         (" in the Specialist Testing column, with detail in the notes: functional testing, soak testing, enhanced "
          "visual and documentation checks", {"color": WHITE})],
        [("Reassess", {"bold": True, "color": LIME}),
         (" if an item once bought new is no longer available from an authorised source: repeat step 1.3 and update "
          "the material master", {"color": WHITE})],
        [("Prefer direct purchase from OEMs; ", {"bold": True, "color": LIME}),
         ("see SEMS 3.3.5.3 for the CFSI process", {"color": WHITE})]]
box(s, 6.76, 1.3, 6.09, 5.15, TEAL, [[("What to do", {"bold": True, "size": 22, "color": LIME})]] + acts,
    size=17, color=WHITE, margin=0.35, space=16)
notes(s, "CFSI is an area where Quality can add real value as intelligent client, because we see supplier "
         "performance and CFSI issues across the business. Use the CFSI / Grey Market factors in Appendix 1. A "
         "CFSI issue found during step 1.6 loops back to step 1.3 for a new risk assessment.")

# ---------------------------------------------------------------- 18 Pitfalls
s = content("1_Comparison", "Common pitfalls and good practice")
pairs = [("Quality writes the QRRA", "Demander leads; Quality facilitates and endorses"),
         ("One grade for the whole system", "Grade at component level, including spares"),
         ("Copying the last QRRA", "Review existing requirements every time"),
         ("Blanket LTRs and inspections", "A clear reason for each activity and record"),
         ("Supporting Comments left blank", "Justify choices and reference Appendix 1 factors"),
         ("Late stakeholder involvement", "ICG, Construction, Commissioning in the room from the start")]
text(s, 0.48, 1.25, 6, 0.4, [[("Pitfall", {"bold": True, "size": 18})]], size=18)
text(s, 6.76, 1.25, 6, 0.4, [[("Good practice", {"bold": True, "size": 18})]], size=18)
for i, (bad, good) in enumerate(pairs):
    y = 1.8 + i * 0.78
    box(s, 0.48, y, 5.85, 0.66, MIST, [bad], size=15, anchor=MSO_ANCHOR.MIDDLE, margin=0.2)
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(6.4), Inches(y + 0.2), Inches(0.3), Inches(0.26))
    a.fill.solid(); a.fill.fore_color.rgb = TEAL; a.line.fill.background()
    box(s, 6.76, y, 6.09, 0.66, PALE, [[(good, {"bold": True})]], size=15, anchor=MSO_ANCHOR.MIDDLE, margin=0.2)
notes(s, "Common issues we see, each paired with what good looks like. Every good-practice point comes straight "
         "from SLP 2.3.2.7_02.")

# ---------------------------------------------------------------- 19 After
s = content("Comparison", "After the session: sign-off and what comes next")
signers = ["Demander", "Project Manager / PEM", "Responsible Engineer / Designer", "Project QE / Quality Manager",
           "Construction Manager / Co-ordinator", "ICG or Inspection Authority", "Commissioning / Pre-operations"]
box(s, 0.48, 1.3, 4.3, 5.15, TEAL, [[("Confirmed by", {"bold": True, "size": 20, "color": LIME})],
                                    [("Delete or add roles as appropriate", {"size": 12, "color": STONE})]] + signers,
    size=15, color=WHITE, margin=0.3, space=12)
nxt = [("Check, approve and issue", "Approve per SLP 1.02.42; issue to Project Document Control per SLP 4.09.108"),
       ("Feed the specification", "Append to the specification with the ITP and LTR schedule; reference on the PR "
        "and material master (SEMS 3.3.7.1); use for tender"),
       ("Manage changes", "Post-award: Contractor Request for Deviation (SEMS 3.3.6.2). Supplier changes: TQ, "
        "concession or production permit (SEMS 2.3.3.1)"),
       ("Feed back", "Off-condition, CFSI or lessons learned go to the Demander, RE and SMEs (step 1.7)")]
for i, (h, b) in enumerate(nxt):
    y = 1.3 + i * 1.32
    badge(s, 5.05, y + 0.25, 0.6, str(i + 1))
    box(s, 5.85, y, 7.0, 1.18, MIST, [[(h, {"bold": True})], b], size=14, margin=0.2, space=2,
        anchor=MSO_ANCHOR.MIDDLE)
notes(s, "From the Guidance tab of SLF 2.3.2.7_02: once fully agreed by all stakeholders, the form is checked and "
         "approved per SLP 1.02.42 and issued to Project Document Control per SLP 4.09.108. The QRRA can be "
         "incorporated into the specification as an appendix, keeping the form format for consistency with ICG "
         "and the supply chain. Post award, the QE is accountable for making sure the oversight and LTR "
         "requirements are adhered to (SEMS 1.2.10.1 and SLP 1.10.315).")

# ---------------------------------------------------------------- 20 References
s = content("1_Comparison", "Reference documents")
left = [["Practices and forms", "Title"],
        ["SLP 2.3.2.7_02", "Determine quality oversight and assurance for procurement of plant and equipment "
                           "(formerly SLP 1.02.45)"],
        ["SLF 2.3.2.7_02", "Quality Requirements Risk Assessment form (formerly SLF 1.02.45.01)"],
        ["SLP 4.06.08", "Procure works, supplies and services"],
        ["SLP 1.02.44", "Prepare and publish a specification"],
        ["SLP 4.06.04", "Quality Assessed Suppliers List (QASL)"],
        ["SLP 4.06.24", "Make vs Buy analysis"],
        ["SLP 1.10.315", "Deliver project quality assurance and control"],
        ["SLSP 1.02.14.02", "Categorisation of nuclear safety, environmental and COMAH consequences"]]
right = [["SEMS", "Title"],
         ["SEMS 2.3.2.7", "Product verification"],
         ["SEMS 2.3.2.6", "Produce the design"],
         ["SEMS 3.3.5.2", "Contract Quality Requirements Manual"],
         ["SEMS 3.3.5.3", "Supplier sub-contractor procurement (CFSI)"],
         ["SEMS 1.2.10.1", "Manage Quality (lifetime records)"],
         ["SEMS 3.3.6.2", "Manage contract change and variations"],
         ["SEMS 2.3.3.1", "Technical queries, concessions, production permits"],
         ["SEMS 3.3.7.1", "Warehousing and inventory (material masters)"]]
table(s, 0.48, 1.3, 6.6, [1.85, 4.75], left, size=12, head_size=13, row_h=0.56)
table(s, 7.3, 1.3, 5.55, [1.75, 3.8], right, size=12, head_size=13, row_h=0.56)
notes(s, "All referenced documents are on SEMS. Remember that SLP 1.02.45 and SLF 1.02.45.01 have been superseded "
         "by SLP 2.3.2.7_02 and SLF 2.3.2.7_02. SLP 1.02.42 (checking and approval) and SLP 4.09.108 (document "
         "control) apply at sign-off.")

# ---------------------------------------------------------------- 21 Close
s = prs.slides.add_slide(L["Title and Content"])
set_text(s.shapes.title, "Key takeaways")
set_text(ph(s, 1), ["The Demander leads and owns the QRRA.",
                    "Quality facilitates, challenges and endorses as the intelligent client.",
                    "Every activity needs a clear, recorded reason."], size=22)
notes(s, "Close by restating the three messages, then take questions. If colleagues are asked to 'do' a QRRA, "
         "the answer is: we will facilitate it with you.")

prs.save(OUT)
print("slides:", len(prs.slides))
